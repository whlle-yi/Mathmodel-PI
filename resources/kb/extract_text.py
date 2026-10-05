"""知识库文本提取脚本：PDF（PyMuPDF）与 DOCX（python-docx）→ 纯文本。

用法：python extract_text.py <文件路径>
输出：UTF-8 纯文本到 stdout（提取失败以非零码退出）。
"""

import sys


def extract_pdf(path: str) -> str:
    import pymupdf

    parts = []
    with pymupdf.open(path) as doc:
        for page in doc:
            parts.append(page.get_text("text"))
    return "\n".join(parts)


def extract_docx(path: str) -> str:
    from docx import Document

    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def main() -> None:
    if len(sys.argv) != 2:
        print("用法: extract_text.py <文件>", file=sys.stderr)
        sys.exit(2)
    path = sys.argv[1].lower()
    try:
        if path.endswith(".pdf"):
            text = extract_pdf(sys.argv[1])
        elif path.endswith(".docx"):
            text = extract_docx(sys.argv[1])
        else:
            with open(sys.argv[1], encoding="utf-8", errors="replace") as f:
                text = f.read()
        sys.stdout.reconfigure(encoding="utf-8")
        print(text)
    except Exception as exc:  # noqa: BLE001
        print(f"提取失败: {exc}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
